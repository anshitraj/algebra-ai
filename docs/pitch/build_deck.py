"""Builds the Algebra 5-minute pitch deck (2 min slides, 3 min demo, close)."""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

OUT = sys.argv[1]

# Brand tokens from frontend/app/algebra.css
PAPER = RGBColor(0xF8, 0xF8, 0xF3)
WASH = RGBColor(0xEE, 0xEF, 0xE7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x22, 0x24, 0x1F)
SOFT = RGBColor(0x66, 0x6B, 0x60)
LINE = RGBColor(0xDE, 0xDF, 0xD6)
SIGNAL = RGBColor(0xBA, 0xD5, 0x9C)
SIGNAL_DIM = RGBColor(0x2D, 0x39, 0x27)
WARN = RGBColor(0xDB, 0xAD, 0x89)
DENY = RGBColor(0xB5, 0x4A, 0x36)

SANS = "Helvetica Neue"
SERIF = "Georgia"
MONO = "Menlo"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
W = 13.333


def bg(slide, color=PAPER):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = color


def text(slide, x, y, w, h, runs, size=18, color=INK, font=SANS, bold=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False, spacing=None):
    """runs: str, or list of paragraphs; each paragraph a str or list of (text, overrides)."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing:
            p.space_after = Pt(spacing)
        for piece in (para if isinstance(para, list) else [(para, {})]):
            t, o = piece if isinstance(piece, tuple) else (piece, {})
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = o.get("font", font)
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", italic)
            f.color.rgb = o.get("color", color)
    return tb


def box(slide, x, y, w, h, fill=WHITE, line=LINE, radius=0.06, weight=1):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.adjustments[0] = radius
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(weight)
    s.shadow.inherit = False
    s.text_frame.text = ""
    return s


def arrow(slide, x1, y1, x2, y2, color=SOFT, weight=1.75):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(weight)
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd",
                          {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    return c


def chrome(slide, n, label):
    """Small mono kicker top-left, page marker bottom-right, hairline."""
    text(slide, 0.7, 0.5, 8, 0.3, label.upper(), size=11, color=SOFT, font=MONO)
    text(slide, W - 2.7, 6.95, 2.0, 0.3, f"ALGEBRA  ·  {n:02d}", size=10, color=SOFT, font=MONO, align=PP_ALIGN.RIGHT)
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.7), Inches(6.8), Inches(W - 0.7), Inches(6.8))
    ln.line.color.rgb = LINE
    ln.line.width = Pt(0.75)


def logo(slide, x, y, s=0.5, color=INK):
    """The interlocking gates: two overlapping rounded squares."""
    for dx in (0, s * 0.5):
        g = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x + dx), Inches(y), Inches(s), Inches(s))
        g.adjustments[0] = 0.28
        g.fill.background()
        g.line.color.rgb = color
        g.line.width = Pt(2.25)


def notes(slide, s):
    slide.notes_slide.notes_text_frame.text = s


# ---------------------------------------------------------------- 1. Title
s = prs.slides.add_slide(BLANK)
bg(s)
logo(s, 0.7, 0.6, 0.42)
text(s, 1.55, 0.66, 4, 0.4, "Algebra", size=22, bold=True)
text(s, 0.7, 2.0, 12, 1.3, "Let your agent pay.", size=72, bold=True)
text(s, 0.7, 3.3, 12, 1.3, "Never let it overspend.", size=68, font=SERIF, italic=True, color=SOFT)
text(s, 0.7, 4.9, 9.5, 1.0,
     "The router and spend firewall for AI agents that pay for APIs in USDC on Solana.",
     size=22, color=INK)
pill = box(s, 0.7, 6.05, 5.6, 0.48, fill=SIGNAL, line=None, radius=0.5)
text(s, 0.7, 6.05, 5.6, 0.48, "●  Live on Solana devnet  ·  program 46gQ…9tdS", size=13, font=MONO,
     color=SIGNAL_DIM, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
notes(s, """[~15s]
Hi, we're building Algebra. AI agents are starting to pay for things themselves: data, search, inference, all paid per call in USDC over x402.
Algebra is the router and the spend firewall that sits between the agent and that money.""")

# ---------------------------------------------------------------- 2. Problem
s = prs.slides.add_slide(BLANK)
bg(s)
chrome(s, 2, "The problem")
text(s, 0.7, 1.0, 12, 1.0, [[("Give an agent a wallet and you get ", {}), ("three problems", {"font": SERIF, "italic": True, "bold": False})]],
     size=36, bold=True)
cards = [
    ("01", "Which provider?",
     "1,600+ paid APIs across four catalogs, under different names and prices. Some are down. Some charge more than they list. Some are traps priced at $25 a call."),
    ("02", "How do I stop it overspending?",
     "“Please be careful” in a prompt is not a budget. You need caps, allow-lists, approvals and a kill switch the model can't talk its way past."),
    ("03", "Did it pay exactly once?",
     "A timeout isn't proof money did or didn't move. Retries and fallbacks must never pay twice."),
]
cw, gap, cy = 3.82, 0.3, 2.45
for i, (n, h, b) in enumerate(cards):
    x = 0.7 + i * (cw + gap)
    box(s, x, cy, cw, 3.85)
    text(s, x + 0.35, cy + 0.35, 1, 0.4, n, size=14, font=MONO, color=SOFT)
    text(s, x + 0.35, cy + 0.85, cw - 0.7, 1.0, h, size=24, bold=True)
    text(s, x + 0.35, cy + 1.95, cw - 0.7, 1.8, b, size=15, color=SOFT)
notes(s, """[~25s]
The moment an agent can pay, three questions come up, and none of them can be left to the model.
One: which provider? There are over 1,600 paid APIs across four catalogs. Some are down, some charge more than they list, and some are honeypots priced at $25 a call, waiting for an agent that pays whatever it's asked.
Two: how do you stop it overspending? A line in the prompt is not a budget.
Three: did it pay exactly once? A timeout doesn't tell you whether money moved.""")

# ---------------------------------------------------------------- 3. How it works
s = prs.slides.add_slide(BLANK)
bg(s)
chrome(s, 3, "How it works")
text(s, 0.7, 1.0, 12, 1.0, [[("Ask for work and a ceiling. ", {}), ("We do the rest.", {"font": SERIF, "italic": True, "bold": False})]],
     size=34, bold=True)
steps = [
    ("Agent", "asks for token.price, at most $0.01", WHITE, INK),
    ("Spend Pass", "budget, per-call cap, allow-list, kill switch", WHITE, INK),
    ("Router", "quotes every provider live; drops down, overpriced and traps", WHITE, INK),
    ("Pay once", "USDC on Solana, from a wallet the agent never sees", WHITE, INK),
    ("Receipt", "the answer and an Ed25519 receipt anyone can verify", SIGNAL, SIGNAL_DIM),
]
bw, bgap, by, bh = 2.18, 0.37, 2.3, 1.75
for i, (h, b, fill, col) in enumerate(steps):
    x = 0.7 + i * (bw + bgap)
    box(s, x, by, bw, bh, fill=fill, line=None if fill == SIGNAL else LINE)
    text(s, x + 0.22, by + 0.25, bw - 0.44, 0.4, h, size=19, bold=True, color=col)
    text(s, x + 0.22, by + 0.8, bw - 0.44, 0.9, b, size=12, color=col if fill == SIGNAL else SOFT)
    if i < len(steps) - 1:
        arrow(s, x + bw + 0.05, by + bh / 2, x + bw + bgap - 0.05, by + bh / 2)
# the call itself
code = box(s, 0.7, 4.45, W - 1.4, 1.95, fill=INK, line=None, radius=0.04)
text(s, 1.05, 4.7, W - 2.1, 1.6, [
    [("POST /api/v1/execute", {"color": SIGNAL, "bold": True}), ("     # REST, or the algebra.execute MCP tool", {"color": SOFT})],
    [('{"capability":"token.price", "budget_max_minor":10000, "provider_policy":{"strategy":"cheapest"}}', {"color": WHITE})],
    [("→ ", {"color": SOFT}), ('delivered · paid demo:beta $0.001 · refused flaky, greedy, trap · receipt eyJhbGci…', {"color": WARN})],
], size=12.5, font=MONO, spacing=10)
notes(s, """[~30s]
Here's how it works. The agent doesn't pick a provider. It asks for a kind of work, say a token price, and the most it will pay.
Its Spend Pass checks the request first: budget, per-call cap, which providers are allowed, how fast it can spend, and a kill switch.
Then the router asks every provider that can do the job for its real price, all at once. It drops the ones that are down, the ones charging more than they list, and the traps.
It pays the best one in USDC on Solana, from a wallet the agent never holds, and hands back the answer with a signed receipt.
It's one REST call, or an MCP tool, so it works with Claude Code, Cursor or the OpenAI Agents SDK.""")

# ---------------------------------------------------------------- 4. On-chain Spend Pass
s = prs.slides.add_slide(BLANK)
bg(s)
chrome(s, 4, "Why Solana")
text(s, 0.7, 1.0, 7.4, 1.6, [[("Limits enforced ", {}), ("on chain,", {"font": SERIF, "italic": True, "bold": False})],
                              [("not just on our server", {})]], size=34, bold=True)
bul = [
    ("Your wallet funds a vault", " that only our Anchor program can move money out of."),
    ("One payment at a time:", " each pull is checked against the per-call cap, budget, window and expiry."),
    ("Freeze it from your wallet.", " The server can't undo that."),
    ("Unspent money goes back", " to the vault, settled from chain proof."),
]
yy = 2.85
for h, b in bul:
    s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.75), Inches(yy + 0.1), Inches(0.13), Inches(0.13)).fill.solid()
    dot = s.shapes[-1]
    dot.fill.fore_color.rgb = SIGNAL_DIM
    dot.line.fill.background()
    text(s, 1.1, yy, 6.6, 0.8, [[(h, {"bold": True}), (b, {"color": SOFT})]], size=17)
    yy += 0.82
# pass card
px, py, pw, ph = 8.45, 1.1, 4.2, 4.6
box(s, px, py, pw, ph, fill=INK, line=None, radius=0.05)
text(s, px + 0.35, py + 0.3, pw - 0.7, 0.3, "SPEND PASS  ·  PDA", size=11, font=MONO, color=SOFT)
text(s, px + 0.35, py + 0.65, pw - 0.7, 0.5, "research-agent", size=22, bold=True, color=WHITE)
rows = [("per-call cap", "$0.05"), ("budget", "$1.00 / month"), ("spent", "$0.003"), ("expires", "in 30 days"), ("frozen", "false")]
ry = py + 1.45
for k, v in rows:
    text(s, px + 0.35, ry, 1.8, 0.35, k, size=13, font=MONO, color=SOFT)
    text(s, px + 2.0, ry, pw - 2.35, 0.35, v, size=13, font=MONO, color=WHITE, align=PP_ALIGN.RIGHT)
    ry += 0.45
box(s, px + 0.35, py + ph - 0.85, pw - 0.7, 0.5, fill=SIGNAL, line=None, radius=0.5)
text(s, px + 0.35, py + ph - 0.85, pw - 0.7, 0.5, "program 46gQ…9tdS · devnet", size=12, font=MONO,
     color=SIGNAL_DIM, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.7, 6.15, 12, 0.5, [[("Even a compromised Algebra server can't take more than the pass allows.", {"font": SERIF, "italic": True})]],
     size=20, color=INK)
notes(s, """[~25s]
And this is why it's on Solana. The Spend Pass isn't only a row in our database. It's an account owned by our own Anchor program, live on devnet.
Your wallet funds a vault. Algebra can pull from it one payment at a time, and the program checks every pull against the cap, the budget, the window and the expiry.
You can freeze it from your own wallet, and we can't undo that. So even if our server were compromised, it couldn't take more than the pass allows.""")

# ---------------------------------------------------------------- 5. Proof
s = prs.slides.add_slide(BLANK)
bg(s)
chrome(s, 5, "Where it stands")
text(s, 0.7, 1.0, 12, 1.0, [[("Built, and ", {}), ("paying on devnet", {"font": SERIF, "italic": True, "bold": False}), (" today", {})]],
     size=36, bold=True)
stats = [
    ("1,622", "providers across Pay.sh, Circle, PayAI and Coinbase's x402 Bazaar"),
    ("14", "kinds of work, one input shape each; 280 endpoints routable today"),
    ("6,131", "paid endpoints tracked, refreshed every six hours, payee pinned"),
    ("24", "tests on the Solana program: 13 unit and 11 LiteSVM"),
]
sw, sgap, sy = 2.85, 0.27, 2.25
for i, (n, l) in enumerate(stats):
    x = 0.7 + i * (sw + sgap)
    box(s, x, sy, sw, 2.3)
    text(s, x + 0.3, sy + 0.3, sw - 0.6, 0.9, n, size=48, bold=True)
    text(s, x + 0.3, sy + 1.25, sw - 0.6, 1.0, l, size=13.5, color=SOFT)
lines = [
    ("Real devnet USDC payments:", " exact, metered through payment channels, and pulled from an on-chain Spend Pass."),
    ("Same request twice, paid once:", " the second call returns the kept answer."),
    ("Mainnet is read-only on purpose:", " quotes and dry runs work; payments wait for an audit."),
]
yy = 4.85
for h, b in lines:
    text(s, 0.7, yy, 12, 0.45, [[("✓  ", {"color": SIGNAL_DIM, "bold": True}), (h, {"bold": True}), (b, {"color": SOFT})]], size=16)
    yy += 0.5
notes(s, """[~20s]
Where it stands: over 1,600 providers from four catalogs, grouped into 14 kinds of work. Real devnet USDC payments, including pulls from the on-chain pass. The same request sent twice pays once.
Mainnet is read-only on purpose until the program is audited.
Now let me show you the core of it: the router and the firewall.""")

# ---------------------------------------------------------------- 6. Demo divider
s = prs.slides.add_slide(BLANK)
bg(s, INK)
text(s, 0.7, 0.5, 8, 0.3, "LIVE DEMO  ·  3 MIN", size=11, color=SOFT, font=MONO)
text(s, 0.7, 1.7, 12, 1.4, [[("Watch it ", {}), ("say no.", {"font": SERIF, "italic": True, "bold": False, "color": SIGNAL})]],
     size=72, bold=True, color=WHITE)
demo = [
    ("1", "Dry run", "four bad providers refused, each with its reason. Nothing paid."),
    ("2", "Execute", "the best one paid once, with a signed receipt."),
    ("3", "Kill switch", "one click and every agent stops, even mid-payment."),
]
yy = 3.65
for n, h, b in demo:
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.7), Inches(yy), Inches(0.5), Inches(0.5))
    c.fill.solid()
    c.fill.fore_color.rgb = SIGNAL
    c.line.fill.background()
    text(s, 0.7, yy, 0.5, 0.5, n, size=16, bold=True, color=SIGNAL_DIM, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 1.5, yy + 0.04, 11, 0.5, [[(h + "  ", {"bold": True, "color": WHITE}), ("— " + b, {"color": LINE})]], size=20,
         anchor=MSO_ANCHOR.MIDDLE)
    yy += 0.85
text(s, 0.7, 6.6, 12, 0.4, "Sandbox: simulated USDC and simulated sellers, some built to misbehave. The same code pays real USDC on devnet.",
     size=12, color=SOFT, font=MONO)
notes(s, """[Leave this up while you switch to the browser]
I'll do this in the sandbox so it's quick and nothing can go wrong with a faucet, but it's the same code path that pays real USDC on devnet.""")

# ---------------------------------------------------------------- 7. Close
s = prs.slides.add_slide(BLANK)
bg(s)
logo(s, 0.7, 0.6, 0.42)
text(s, 1.55, 0.66, 4, 0.4, "Algebra", size=22, bold=True)
text(s, 0.7, 1.9, 12, 1.2, "Agents can spend.", size=64, bold=True)
text(s, 0.7, 3.05, 12, 1.2, "Algebra decides how much.", size=60, font=SERIF, italic=True, color=SOFT)
nxt = [("Next", "Program audit, then mainnet payments behind on-chain passes"),
       ("Next", "More catalogs and kinds of work; paid plans for teams"),
       ("Try it", "REST or MCP: one key, a Spend Pass, and your agent is ready")]
yy = 4.5
for k, v in nxt:
    text(s, 0.7, yy, 1.3, 0.4, k.upper(), size=12, font=MONO, color=SOFT)
    text(s, 2.0, yy - 0.03, 10, 0.45, v, size=18)
    yy += 0.55
text(s, 0.7, 6.3, 7, 0.4, "github.com/anshitraj/algebra-ai", size=14, font=MONO, color=INK)
text(s, W - 4.2, 6.15, 3.5, 0.6, "Thank you.", size=30, font=SERIF, italic=True, align=PP_ALIGN.RIGHT)
notes(s, """[~15s]
So that's Algebra: agents can spend, and Algebra decides how much.
Next up is an audit of the program, then real mainnet payments behind on-chain passes.
Thank you, happy to take questions.""")

prs.save(OUT)
print("saved", OUT)
